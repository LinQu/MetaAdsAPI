"""Penanaman gambar iklan pada workbook XLSX."""

import hashlib
import io
import os
from typing import Any, Dict, List, Sequence, Tuple

import requests
from openpyxl.drawing.image import Image as ExcelImage

from ...logging_utils import warning

try:
    from PIL import Image as PILImage
except ImportError:  # Image embedding is optional
    PILImage = None  # type: ignore


def download_image_to_png(
    url: str,
    temp_directory: str,
    width: int,
    height: int,
    image_session: requests.Session,
) -> Tuple[str, int, int]:
    if not url.startswith(("http://", "https://")):
        raise ValueError("URL gambar tidak valid.")

    response = image_session.get(url, timeout=(10, 60), stream=True)
    response.raise_for_status()
    content = response.content
    if len(content) > 15 * 1024 * 1024:
        raise ValueError("Ukuran gambar melebihi 15 MB.")

    if PILImage is None:
        raise RuntimeError("Pillow belum terpasang; gunakan --tanpa-gambar atau install Pillow.")
    image = PILImage.open(io.BytesIO(content))
    image.load()
    if image.mode not in ("RGB", "RGBA"):
        image = image.convert("RGBA")
    image.thumbnail((width, height))

    digest = hashlib.sha256(url.encode("utf-8")).hexdigest()[:20]
    path = os.path.join(temp_directory, "{}.png".format(digest))
    image.save(path, format="PNG")
    return path, int(image.width), int(image.height)


def embed_rekap_images(
    worksheet: Any,
    rows: List[Dict[str, Any]],
    columns: Sequence[str],
    errors: List[Dict[str, Any]],
    temp_directory: str,
    width: int,
    height: int,
) -> None:
    image_column = columns.index("Gambar iklan") + 1
    url_column = columns.index("URL gambar") + 1
    image_session = requests.Session()
    image_session.headers.update({"User-Agent": "MetaAdsReportImageFetcher/1.0"})
    cache: Dict[str, Tuple[str, int, int]] = {}

    for offset, row in enumerate(rows, start=2):
        url = str(row.get("URL gambar", "")).strip()
        url_cell = worksheet.cell(offset, url_column)
        if url:
            url_cell.hyperlink = url
            url_cell.style = "Hyperlink"
        if not url:
            worksheet.cell(offset, image_column).value = "Tidak tersedia"
            continue

        try:
            if url not in cache:
                cache[url] = download_image_to_png(
                    url,
                    temp_directory,
                    width,
                    height,
                    image_session,
                )
            image_path, image_width, image_height = cache[url]
            excel_image = ExcelImage(image_path)
            excel_image.width = image_width
            excel_image.height = image_height
            anchor_cell = worksheet.cell(offset, image_column).coordinate
            worksheet.add_image(excel_image, anchor_cell)
            worksheet.row_dimensions[offset].height = max(78, image_height * 0.78)
        except Exception as exc:
            worksheet.cell(offset, image_column).value = "Gagal dimuat"
            errors.append({
                "ID ACCOUNT": str(row.get("ID ACCOUNT", "")),
                "CABANG": str(row.get("CABANG", "")),
                "BISNIS": str(row.get("BISNIS", "")),
                "ID iklan": str(row.get("ID iklan", "")),
                "Tahap": "Gambar",
                "Error": str(exc),
            })
            warning(
                "Gambar untuk iklan {} gagal dimuat: {}".format(
                    row.get("ID iklan", ""),
                    exc,
                )
            )
