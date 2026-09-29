"""Ekspor workbook XLSX lengkap dengan sheet Rekap/Rinci, Error, Informasi."""

import os
import tempfile
from typing import Any, Dict, List

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ...constants import ERROR_COLUMNS
from ..info import information_rows
from ..result import ExportResult
from ..schema import report_columns
from ..values import excel_value, sanitize_spreadsheet_text
from .images import PILImage, embed_rekap_images
from .styles import (
    apply_number_formats,
    set_report_widths,
    style_error_sheet,
    style_header,
)


def export_xlsx(
    output_file: str,
    mode: str,
    rows: List[Dict[str, Any]],
    errors: List[Dict[str, Any]],
    input_file: str,
    api_version: str,
    since: str,
    until: str,
    status_filter: str,
    date_source: str,
    image_width: int,
    image_height: int,
    embed_images: bool,
    input_source: str = "excel",
    source_reference: str = "",
) -> ExportResult:
    if embed_images and PILImage is None:
        raise RuntimeError("Pillow belum terpasang; gunakan --tanpa-gambar atau install Pillow.")
    columns = report_columns(mode)
    workbook = Workbook()
    report_sheet = workbook.active
    report_sheet.title = "Rekap Iklan" if mode == "rekap" else "Rinci Iklan"
    report_sheet.append(columns)
    for row in rows:
        report_sheet.append([
            excel_value(column_name, row.get(column_name, ""))
            for column_name in columns
        ])

    style_header(report_sheet, columns, len(rows))
    set_report_widths(report_sheet, columns)
    apply_number_formats(report_sheet, columns, len(rows))

    output_directory = os.path.dirname(os.path.abspath(output_file))
    if output_directory:
        os.makedirs(output_directory, exist_ok=True)

    with tempfile.TemporaryDirectory(prefix="meta_ads_images_") as temp_directory:
        if mode == "rekap" and embed_images:
            embed_rekap_images(
                report_sheet,
                rows,
                columns,
                errors,
                temp_directory,
                image_width,
                image_height,
            )

        error_sheet = workbook.create_sheet("Error")
        error_sheet.append(ERROR_COLUMNS)
        for item in errors:
            error_sheet.append([
                sanitize_spreadsheet_text(item.get(column, ""))
                for column in ERROR_COLUMNS
            ])
        style_error_sheet(error_sheet, len(errors))

        info_sheet = workbook.create_sheet("Informasi")
        info_sheet.sheet_view.showGridLines = False
        info_rows = information_rows(
            mode=mode,
            status_filter=status_filter,
            api_version=api_version,
            since=since,
            until=until,
            input_file=input_file,
            date_source=date_source,
            output_format="xlsx",
            images_enabled=mode == "rekap" and embed_images,
            input_source=input_source,
            source_reference=source_reference,
        )
        for row in info_rows:
            info_sheet.append([sanitize_spreadsheet_text(value) for value in row])
        info_sheet.column_dimensions["A"].width = 26
        info_sheet.column_dimensions["B"].width = 100
        for cell in info_sheet[1]:
            cell.fill = PatternFill("solid", fgColor="7030A0")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")

        workbook.save(output_file)
    workbook.close()

    # Structural verification: ensure the generated workbook can be reopened.
    verification = load_workbook(output_file, read_only=True, data_only=False)
    try:
        expected = {report_sheet.title, "Error", "Informasi"}
        if not expected.issubset(set(verification.sheetnames)):
            raise RuntimeError("Verifikasi XLSX gagal: sheet output tidak lengkap.")
    finally:
        verification.close()

    return ExportResult(report_file=os.path.abspath(output_file))
