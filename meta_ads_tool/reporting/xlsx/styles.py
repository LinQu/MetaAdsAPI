"""Styling worksheet XLSX."""

from typing import Any, Sequence

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..schema import (
    _DATE_COLUMNS,
    _DATETIME_COLUMNS,
    _ID_TEXT_COLUMNS,
)


def style_header(worksheet: Any, columns: Sequence[str], row_count: int) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    bottom_border = Border(bottom=Side(style="thin", color="A6A6A6"))

    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 36
    worksheet.auto_filter.ref = "A1:{}{}".format(
        get_column_letter(len(columns)),
        max(row_count + 1, 1),
    )
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = bottom_border


def set_report_widths(worksheet: Any, columns: Sequence[str]) -> None:
    widths_by_name = {
        "ID ACCOUNT": 22,
        "Nama akun": 28,
        "ID kampanye": 22,
        "Nama kampanye": 34,
        "Tujuan kampanye": 22,
        "ID set iklan": 22,
        "Nama set iklan": 34,
        "Optimasi": 24,
        "ID iklan": 22,
        "Nama iklan": 34,
        "Status efektif": 18,
        "Status konfigurasi": 18,
        "Gambar iklan": 25,
        "URL gambar": 42,
        "Tanggal": 13,
        "Waktu (zona waktu akun iklan)": 24,
        "Jam mulai": 12,
        "Kategori waktu": 16,
        "Pengaturan atribusi": 24,
        "Mulai": 24,
        "Berakhir": 24,
        "Impresi": 14,
        "Jangkauan": 14,
        "Frekuensi": 13,
        "Klik tautan": 14,
        "CTR klik tautan": 17,
        "Biaya per klik tautan": 20,
        "Hasil": 14,
        "Biaya per hasil": 18,
        "Jumlah yang dibelanjakan": 24,
        "Mata uang": 12,
        "Awal pelaporan": 16,
        "Akhir pelaporan": 16,
        "CABANG": 18,
        "BISNIS": 20,
    }
    for index, column_name in enumerate(columns, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = widths_by_name.get(
            column_name,
            18,
        )


def apply_number_formats(
    worksheet: Any,
    columns: Sequence[str],
    row_count: int,
) -> None:
    column_index = {name: index + 1 for index, name in enumerate(columns)}

    for row_number in range(2, row_count + 2):
        for name in _ID_TEXT_COLUMNS:
            if name in column_index:
                worksheet.cell(row_number, column_index[name]).number_format = "@"
        for name in _DATE_COLUMNS:
            if name in column_index:
                worksheet.cell(row_number, column_index[name]).number_format = "yyyy-mm-dd"
        for name in _DATETIME_COLUMNS:
            if name in column_index:
                worksheet.cell(row_number, column_index[name]).number_format = (
                    "yyyy-mm-dd hh:mm:ss"
                )
        for name in ("Impresi", "Jangkauan", "Klik tautan", "Jam mulai"):
            if name in column_index:
                worksheet.cell(row_number, column_index[name]).number_format = "#,##0"
        for name in (
            "Frekuensi",
            "Biaya per klik tautan",
            "Hasil",
            "Biaya per hasil",
            "Jumlah yang dibelanjakan",
        ):
            if name in column_index:
                worksheet.cell(row_number, column_index[name]).number_format = "#,##0.00"
        if "CTR klik tautan" in column_index:
            worksheet.cell(
                row_number,
                column_index["CTR klik tautan"],
            ).number_format = "0.00%"
        for column_number in range(1, len(columns) + 1):
            worksheet.cell(row_number, column_number).alignment = Alignment(
                vertical="top",
                wrap_text=False,
            )


def style_error_sheet(worksheet: Any, error_count: int) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 28
    for cell in worksheet[1]:
        cell.fill = PatternFill("solid", fgColor="C00000")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.column_dimensions["A"].width = 22
    worksheet.column_dimensions["B"].width = 18
    worksheet.column_dimensions["C"].width = 20
    worksheet.column_dimensions["D"].width = 22
    worksheet.column_dimensions["E"].width = 22
    worksheet.column_dimensions["F"].width = 100
    if error_count:
        worksheet.auto_filter.ref = "A1:F{}".format(error_count + 1)
